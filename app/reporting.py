from __future__ import annotations

import json
import os
from collections import Counter, defaultdict
from contextlib import closing
from datetime import datetime
from io import BytesIO
from pathlib import Path
from typing import Any

from fastapi import HTTPException
from dotenv import load_dotenv
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .questions import QUESTIONS

load_dotenv(Path(__file__).resolve().parents[1] / ".env")
PASSING_PERCENT = int(os.getenv("PASSING_PERCENT", "80"))


def collect_report(connect_db, date_from: str | None = None, date_to: str | None = None, name: str | None = None) -> dict[str, Any]:
    for label, value in (("inicial", date_from), ("final", date_to)):
        if value:
            try:
                datetime.strptime(value, "%Y-%m-%d")
            except ValueError as exc:
                raise HTTPException(status_code=422, detail=f"La fecha {label} no es válida.") from exc
    if date_from and date_to and date_from > date_to:
        raise HTTPException(status_code=422, detail="La fecha inicial no puede ser posterior a la final.")
    conditions, params = [], []
    if date_from:
        conditions.append("date(submitted_at) >= date(?)"); params.append(date_from)
    if date_to:
        conditions.append("date(submitted_at) <= date(?)"); params.append(date_to)
    if name and name.strip():
        conditions.append("participant_name LIKE ? COLLATE NOCASE"); params.append(f"%{name.strip()}%")
    where = f"WHERE {' AND '.join(conditions)}" if conditions else ""
    with closing(connect_db()) as db:
        rows = db.execute(f"SELECT id, participant_name, score, total, submitted_at, answers_json FROM attempts {where} ORDER BY id DESC", params).fetchall()

    results, answer_rows = [], []
    section_counts = defaultdict(lambda: [0, 0])
    question_counts = {q["id"]: {"question_id": q["id"], "section": q["section"], "question": q["question"], "correct": 0, "incorrect": 0, "wrong": Counter()} for q in QUESTIONS}
    for row in rows:
        result = {key: row[key] for key in ("id", "participant_name", "score", "total", "submitted_at")}
        result["percent"] = round(result["score"] / result["total"] * 100) if result["total"] else 0
        result["passed"] = result["percent"] >= PASSING_PERCENT
        results.append(result)
        for answer in json.loads(row["answers_json"]):
            section_counts[answer["section"]][1] += 1
            q = question_counts[answer["question_id"]]
            if answer["is_correct"]:
                section_counts[answer["section"]][0] += 1; q["correct"] += 1
            else:
                q["incorrect"] += 1; q["wrong"][answer["selected"]] += 1
            answer_rows.append({"attempt_id": row["id"], "participant_name": row["participant_name"], "submitted_at": row["submitted_at"], "section": answer["section"], "question_id": answer["question_id"], "question": answer["question"], "selected": answer["selected"], "correct_answer": answer["correct"], "is_correct": answer["is_correct"]})
    count = len(results)
    passed = sum(r["passed"] for r in results)
    sections = [{"section": s, "correct": v[0], "total": v[1], "percent": round(v[0] / v[1] * 100, 1) if v[1] else 0} for s, v in sorted(section_counts.items())]
    questions = []
    for q in question_counts.values():
        attempts = q["correct"] + q["incorrect"]
        wrong = q.pop("wrong")
        q.update(attempts=attempts, percent_correct=round(q["correct"] / attempts * 100, 1) if attempts else 0, most_common_wrong=wrong.most_common(1)[0][0] if wrong else "")
        questions.append(q)
    questions.sort(key=lambda q: (q["percent_correct"], q["question_id"]))
    return {"total_attempts": count, "average_score": round(sum(r["score"] for r in results) / count, 1) if count else 0, "average_percent": round(sum(r["percent"] for r in results) / count, 1) if count else 0, "passing_percent": PASSING_PERCENT, "passed_count": passed, "needs_refuerzo_count": count - passed, "pass_rate": round(passed / count * 100, 1) if count else 0, "date_from": date_from or "", "date_to": date_to or "", "name_filter": name or "", "results": results, "sections": sections, "questions": questions, "answer_rows": answer_rows}


def make_workbook(report: dict[str, Any]) -> BytesIO:
    wb = Workbook()
    ws = wb.active; ws.title = "Resumen"
    ws.append(["Reporte de evaluación · Sistema de Gestión de la Calidad"])
    for label, value in [("Periodo desde", report["date_from"] or "Sin límite"), ("Periodo hasta", report["date_to"] or "Sin límite"), ("Filtro por nombre", report["name_filter"] or "Todos"), ("Umbral de aprobación", f'{report["passing_percent"]}%')]: ws.append([label, value])
    ws.append([]); ws.append(["Indicador", "Resultado"])
    for label, value in [("Evaluaciones completadas", report["total_attempts"]), ("Calificación promedio", f'{report["average_score"]} / {len(QUESTIONS)}'), ("Promedio porcentual", f'{report["average_percent"]}%'), ("Aprobadas", report["passed_count"]), ("Requieren refuerzo", report["needs_refuerzo_count"]), ("Porcentaje de aprobación", f'{report["pass_rate"]}%')]: ws.append([label, value])
    ws = wb.create_sheet("Por sección"); ws.append(["Sección", "Aciertos", "Reactivos respondidos", "% de aciertos"])
    for q in report["sections"]: ws.append([q["section"], q["correct"], q["total"], q["percent"] / 100])
    ws = wb.create_sheet("Por pregunta"); ws.append(["N.º", "Sección", "Pregunta", "Evaluadas", "Correctas", "Incorrectas", "% correctas", "Respuesta incorrecta más frecuente"])
    for q in report["questions"]: ws.append([q["question_id"], q["section"], q["question"], q["attempts"], q["correct"], q["incorrect"], q["percent_correct"] / 100, q["most_common_wrong"]])
    ws = wb.create_sheet("Resultados individuales"); ws.append(["Folio", "Nombre", "Fecha y hora", "Aciertos", "Total", "Porcentaje", "Estatus"])
    for r in report["results"]: ws.append([r["id"], r["participant_name"], r["submitted_at"], r["score"], r["total"], r["percent"] / 100, "Aprobado" if r["passed"] else "Requiere refuerzo"])
    ws = wb.create_sheet("Respuestas"); ws.append(["Folio", "Nombre", "Fecha y hora", "Sección", "N.º pregunta", "Pregunta", "Respuesta elegida", "Respuesta correcta", "Resultado"])
    for a in report["answer_rows"]: ws.append([a["attempt_id"], a["participant_name"], a["submitted_at"], a["section"], a["question_id"], a["question"], a["selected"], a["correct_answer"], "Correcta" if a["is_correct"] else "Incorrecta"])
    dark = PatternFill("solid", fgColor="0F172A")
    for sheet in wb.worksheets:
        sheet.freeze_panes = "A7" if sheet.title == "Resumen" else "A2"
        if sheet.title != "Resumen":
            sheet.auto_filter.ref = sheet.dimensions
            for cell in sheet[1]: cell.fill = dark; cell.font = Font(color="FFFFFF", bold=True)
        else:
            sheet[1][0].font = Font(size=16, bold=True, color="0F172A")
            for cell in sheet[7]: cell.fill = dark; cell.font = Font(color="FFFFFF", bold=True)
        for column in sheet.columns:
            idx = column[0].column
            sheet.column_dimensions[get_column_letter(idx)].width = min(max(max((len(str(c.value or "")) for c in column[:1000]), default=12) + 2, 14), 58)
        for row in sheet.iter_rows():
            for cell in row: cell.alignment = Alignment(vertical="top", wrap_text=True)
    for sheet_name, col in [("Por sección", 4), ("Por pregunta", 7), ("Resultados individuales", 6)]:
        for cells in wb[sheet_name].iter_cols(min_col=col, max_col=col, min_row=2):
            for cell in cells: cell.number_format = "0.0%"
    output = BytesIO(); wb.save(output); output.seek(0)
    return output
